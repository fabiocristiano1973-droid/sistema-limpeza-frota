import openpyxl, json, pickle, statistics, datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from collections import defaultdict

# >>> ATUALIZE aqui para apontar para o mesmo arquivo de boletim usado no script 01 <<<
SRC = '../planilha/Boletim_do_Veiculo_mais_recente.xlsx'
OUT = 'Gestao_Media_Combustivel_ROTA_Itabuna.xlsx'

d = pickle.load(open('state.pkl', 'rb'))
res = json.load(open('resultado.json'))
agg = json.load(open('agg_vehicle.json'))

rows = d['rows']; groups = d['groups']; matched = d['matched']; unmatched = d['unmatched']
frota_rows = d['frota_rows']

FONT_NAME = 'Arial'
BLUE = Font(name=FONT_NAME, color='0000FF')
BLUE_BOLD = Font(name=FONT_NAME, color='0000FF', bold=True)
BLACK = Font(name=FONT_NAME, color='000000')
GREEN = Font(name=FONT_NAME, color='008000')
HEADER_FONT = Font(name=FONT_NAME, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='1F3864')
SUBHEADER_FILL = PatternFill('solid', fgColor='D9E1F2')
YELLOW_FILL = PatternFill('solid', fgColor='FFFF00')
RED_FILL = PatternFill('solid', fgColor='FFC7CE')
GREEN_FILL = PatternFill('solid', fgColor='C6EFCE')
YELLOW_FILL2 = PatternFill('solid', fgColor='FFEB9C')
GRAY_FILL = PatternFill('solid', fgColor='F2F2F2')
TITLE_FONT = Font(name=FONT_NAME, bold=True, size=14, color='1F3864')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def style_header(ws, row, ncols, start_col=1):
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = BORDER

def autofit(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

def set_font_all(ws, font=BLACK):
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and cell.font.name != FONT_NAME:
                cell.font = Font(name=FONT_NAME, size=cell.font.size or 10,
                                  bold=cell.font.bold, color=cell.font.color)

print("Loading source workbook...")
wb = openpyxl.load_workbook(SRC)
ws_bol = wb.worksheets[0]  # aba do boletim = sempre a primeira do arquivo exportado
ws_frota = wb['Base_Frota']

# ============================================================
# 1) Helper columns appended to the original Boletim sheet
# ============================================================
print("Adding helper columns to Boletim sheet...")
last_col = ws_bol.max_column  # 63
helper_start = last_col + 2  # leave one blank col as visual separator (64), helpers start 65
sep_col = last_col + 1
h1, h2, h3, h4, h5 = helper_start, helper_start+1, helper_start+2, helper_start+3, helper_start+4

headers = {
    h1: 'Distância (km) - todos os dias',
    h2: 'Consumo (L) - dias com reporte',
    h3: 'Média (Km/L) - dias com reporte',
    h4: 'Distância (km) - dias com reporte',
    h5: 'Tracinho (1=sem reporte telemetria)',
}
for col, text in headers.items():
    cell = ws_bol.cell(row=1, column=col, value=text)
    cell.font = HEADER_FONT
    cell.fill = PatternFill('solid', fgColor='548235')
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
note = ws_bol.cell(row=1, column=sep_col)
note.fill = PatternFill('solid', fgColor='FFFFFF')

def locale_safe_number(expr):
    # Converte um texto numerico decimal com ponto (ex. "126.33") em numero
    # SEM usar VALUE()/NUMBERVALUE() - em Excel configurado em pt-BR (virgula
    # decimal), VALUE("126.33") falha (#VALUE!) porque o ponto nao e separador
    # decimal valido nesse idioma. Aqui a parte inteira e a fracionaria sao
    # extraidas como substrings de digitos puros (sem separador nenhum) e
    # multiplicadas/somadas — a conversao texto->numero de um digito puro e
    # identica em qualquer configuracao regional do Excel.
    return (f'(LEFT({expr},FIND(".",{expr})-1)*1+'
            f'MID({expr},FIND(".",{expr})+1,LEN({expr})-FIND(".",{expr}))/'
            f'POWER(10,LEN({expr})-FIND(".",{expr})))')

max_row = ws_bol.max_row
for r in range(2, max_row + 1):
    ws_bol.cell(row=r, column=h1, value=f'=IFERROR({locale_safe_number(f"F{r}")},"")')
    ws_bol.cell(row=r, column=h2, value=f'=IF(R{r}="-","",IFERROR({locale_safe_number(f"LEFT(Q{r},LEN(Q{r})-2)")},""))')
    ws_bol.cell(row=r, column=h3, value=f'=IF(R{r}="-","",IFERROR({locale_safe_number(f"LEFT(R{r},LEN(R{r})-5)")},""))')
    ws_bol.cell(row=r, column=h4, value=f'=IF(R{r}="-","",IFERROR({locale_safe_number(f"F{r}")},""))')
    ws_bol.cell(row=r, column=h5, value=f'=IF(R{r}="-",1,0)')
    for col in (h1, h2, h3, h4):
        ws_bol.cell(row=r, column=col).number_format = '0.00'

ws_bol.column_dimensions[get_column_letter(sep_col)].width = 2
for col in (h1, h2, h3, h4, h5):
    ws_bol.column_dimensions[get_column_letter(col)].width = 16

HCOL = {'dist_total': get_column_letter(h1), 'litros_rep': get_column_letter(h2),
        'media_rep': get_column_letter(h3), 'dist_rep': get_column_letter(h4),
        'tracinho': get_column_letter(h5)}
BOL = ws_bol.title
BOLQ = f"'{BOL}'"  # quoted sheet reference for formulas (name contains hyphens)
print("Helper columns added at:", HCOL, "-> sheet", BOL)

# ============================================================
# 2) METAS sheet
# ============================================================
print("Building Metas sheet...")
ws_meta = wb.create_sheet('Metas')
ws_meta.sheet_view.showGridLines = False

ws_meta['B2'] = 'METAS E PARÂMETROS — GESTÃO DE MÉDIA DE COMBUSTÍVEL'
ws_meta['B2'].font = TITLE_FONT
ws_meta.merge_cells('B2:G2')

ws_meta['B4'] = 'Preço médio do diesel (R$/L)'
ws_meta['B4'].font = Font(name=FONT_NAME, bold=True)
ws_meta['D4'] = 6.00
ws_meta['D4'].font = BLUE_BOLD
ws_meta['D4'].fill = YELLOW_FILL
ws_meta['D4'].number_format = 'R$ #,##0.00'
ws_meta['D4'].border = BORDER
ws_meta['D4'].comment = Comment(
    "PREENCHER: a base de dados do boletim de telemetria não informa o preço pago pelo diesel. "
    "Valor de R$ 6,00/L inserido como referência de mercado a ser SUBSTITUÍDA pelo preço médio "
    "real pago pela empresa no período (nota fiscal / contrato do fornecedor).", "Auditoria")
ws_meta['E4'] = '<< ATUALIZAR com o preço real pago (não fornecido na base) >>'
ws_meta['E4'].font = Font(name=FONT_NAME, italic=True, color='C00000', size=9)

ws_meta['B6'] = 'Metas de eficiência por Modelo/Fabricante × Unidade Operacional (Setor)'
ws_meta['B6'].font = Font(name=FONT_NAME, bold=True, size=11)
ws_meta['B7'] = ('Meta inicial sugerida = mediana do km/L real dos veículos válidos do próprio grupo, apurada nesta base '
                  '(08/09 a 14/09/2026). Ajuste manualmente para a meta oficial da empresa quando disponível.')
ws_meta['B7'].font = Font(name=FONT_NAME, italic=True, size=9, color='595959')
ws_meta.merge_cells('B7:H7')

headers_meta = ['Fabricante/Modelo (FB)', 'Unidade Operacional (Setor)', 'Nº de veículos na Base_Frota',
                'Km/L real do grupo no período (agregado)', 'Meta (Km/L) — editável', 'Fonte da meta']
r0 = 9
for i, h in enumerate(headers_meta):
    ws_meta.cell(row=r0, column=2 + i, value=h)
style_header(ws_meta, r0, len(headers_meta), start_col=2)

by_group = defaultdict(lambda: {'km': 0.0, 'l': 0.0, 'kmls': [], 'n': 0})
for pref, a in agg.items():
    g = (a['fb'], a['setor'])
    by_group[g]['km'] += a['km_reportado']
    by_group[g]['l'] += a['litros_reportado']
    by_group[g]['n'] += 1
    if a['kml']:
        by_group[g]['kmls'].append(a['kml'])

frota_group_count = defaultdict(int)
for f in frota_rows:
    frota_group_count[(f['fb'], f['setor'])] += 1

meta_rows_map = {}
r = r0 + 1
for g in sorted(frota_group_count.keys()):
    fb, setor = g
    v = by_group.get(g, {'km': 0, 'l': 0, 'kmls': [], 'n': 0})
    kml_group = round(v['km'] / v['l'], 2) if v['l'] else None
    median_kml = round(statistics.median(v['kmls']), 2) if v['kmls'] else None
    ws_meta.cell(row=r, column=2, value=fb).font = BLACK
    ws_meta.cell(row=r, column=3, value=setor).font = BLACK
    ws_meta.cell(row=r, column=4, value=frota_group_count[g]).font = BLACK
    c_kml = ws_meta.cell(row=r, column=5, value=kml_group if kml_group else 'Sem dados no período')
    c_kml.font = BLACK
    if kml_group:
        c_kml.number_format = '0.00'
    meta_cell = ws_meta.cell(row=r, column=6,
                              value=median_kml if median_kml else (kml_group if kml_group else ''))
    meta_cell.font = BLUE_BOLD
    meta_cell.fill = YELLOW_FILL2
    meta_cell.number_format = '0.00'
    meta_cell.border = BORDER
    ws_meta.cell(row=r, column=7,
                 value='Mediana do grupo (sugestão inicial) — substituir por meta oficial' if median_kml
                 else 'Sem veículos com reporte numérico no período')
    ws_meta.cell(row=r, column=7).font = Font(name=FONT_NAME, italic=True, size=9, color='595959')
    for c in range(2, 8):
        ws_meta.cell(row=r, column=c).border = BORDER
    meta_rows_map[g] = r
    r += 1
meta_last_row = r - 1

autofit(ws_meta, [3, 20, 22, 20, 22, 16, 40])
ws_meta.freeze_panes = 'B10'

# ============================================================
# 3) ANALISE_CONSUMO sheet (per-vehicle km/L, meta, deviation, cost, savings)
# ============================================================
print("Building Analise_Consumo sheet...")
ws_ac = wb.create_sheet('Analise_Consumo')
ws_ac.sheet_view.showGridLines = False
ws_ac['A1'] = 'GESTÃO DE MÉDIA DE COMBUSTÍVEL — ANÁLISE POR VEÍCULO (km/L real = km ÷ litros, não média das médias)'
ws_ac['A1'].font = TITLE_FONT
ws_ac.merge_cells('A1:S1')
ws_ac['A2'] = ('km/L real calculado apenas com os dias em que a telemetria reportou número (média ≠ "-"), '
               'para não distorcer o indicador com dias sem medição. Base: Boletim 08/09 a 14/09/2026.')
ws_ac['A2'].font = Font(name=FONT_NAME, italic=True, size=9, color='595959')
ws_ac.merge_cells('A2:S2')

ac_headers = ['Prefixo', 'Placa', 'Fabricante/Modelo', 'Setor/UO', 'Classificação de reporte',
              'Km totais (todos os dias)', 'Km em dias com reporte', 'Litros em dias com reporte',
              'Dias com reporte', 'Dias com tracinho', 'Km/L real', 'Meta (Km/L)', 'Desvio vs meta (%)',
              'Status', 'Preço diesel (R$/L)', 'Custo por Km (R$)', 'Custo total no período (R$)',
              'Economia potencial (L)', 'Economia potencial (R$)']
r0 = 4
for i, h in enumerate(ac_headers):
    ws_ac.cell(row=r0, column=1 + i, value=h)
style_header(ws_ac, r0, len(ac_headers))
ws_ac.freeze_panes = 'A5'

class_by_prefixo = {a['prefixo']: a['classificacao'] for a in res['analise']}
fb_by_prefixo = {a['prefixo']: a['fb'] for a in res['analise']}
setor_by_prefixo = {a['prefixo']: a['setor'] for a in res['analise']}

r = r0 + 1
first_data_row = r
sorted_prefixos = sorted(agg.keys())
for pref in sorted_prefixos:
    fb = fb_by_prefixo.get(pref, '')
    setor = setor_by_prefixo.get(pref, '')
    placa_match = [a['placa'] for a in res['analise'] if a['prefixo'] == pref]
    placa = placa_match[0] if placa_match else ''
    classif = class_by_prefixo.get(pref, '')

    ws_ac.cell(row=r, column=1, value=pref)
    ws_ac.cell(row=r, column=2, value=placa)
    ws_ac.cell(row=r, column=3, value=fb)
    ws_ac.cell(row=r, column=4, value=setor)
    ws_ac.cell(row=r, column=5, value=classif)

    kmtot = f'=SUMIFS({BOLQ}!${HCOL["dist_total"]}:${HCOL["dist_total"]},{BOLQ}!$B:$B,$A{r})'
    kmrep = f'=SUMIFS({BOLQ}!${HCOL["dist_rep"]}:${HCOL["dist_rep"]},{BOLQ}!$B:$B,$A{r})'
    litrep = f'=SUMIFS({BOLQ}!${HCOL["litros_rep"]}:${HCOL["litros_rep"]},{BOLQ}!$B:$B,$A{r})'
    diasrep = f'=COUNTIFS({BOLQ}!$B:$B,$A{r},{BOLQ}!$R:$R,"<>-")'
    diastrac = f'=COUNTIFS({BOLQ}!$B:$B,$A{r},{BOLQ}!$R:$R,"-")'
    ws_ac.cell(row=r, column=6, value=kmtot)
    ws_ac.cell(row=r, column=7, value=kmrep)
    ws_ac.cell(row=r, column=8, value=litrep)
    ws_ac.cell(row=r, column=9, value=diasrep)
    ws_ac.cell(row=r, column=10, value=diastrac)

    kml = f'=IF(H{r}=0,"Sem dados",G{r}/H{r})'
    ws_ac.cell(row=r, column=11, value=kml)

    meta_lookup = (f'=IFERROR(INDEX(Metas!$F$10:$F${meta_last_row},'
                   f'MATCH(1,INDEX((Metas!$B$10:$B${meta_last_row}=C{r})*(Metas!$C$10:$C${meta_last_row}=D{r}),0),0)),"")')
    ws_ac.cell(row=r, column=12, value=meta_lookup)

    desvio = f'=IF(OR(K{r}="Sem dados",L{r}="",L{r}=0),"",(K{r}-L{r})/L{r})'
    ws_ac.cell(row=r, column=13, value=desvio)
    ws_ac.cell(row=r, column=13).number_format = '0.0%'

    status = (f'=IF(K{r}="Sem dados","Sem dados suficientes",IF(M{r}="","Sem meta definida",'
              f'IF(M{r}>=0.02,"Acima da meta",IF(M{r}<=-0.10,"Crítico: abaixo da meta",'
              f'IF(M{r}<0,"Abaixo da meta","Na meta")))))')
    ws_ac.cell(row=r, column=14, value=status)

    ws_ac.cell(row=r, column=15, value='=Metas!$D$4')
    ws_ac.cell(row=r, column=15).number_format = 'R$ #,##0.00'

    custokm = f'=IF(K{r}="Sem dados","",IFERROR(O{r}/K{r},"Km=0 c/ consumo>0 (ver nota)"))'
    ws_ac.cell(row=r, column=16, value=custokm)
    ws_ac.cell(row=r, column=16).number_format = 'R$ #,##0.00'

    custotot = f'=H{r}*O{r}'
    ws_ac.cell(row=r, column=17, value=custotot)
    ws_ac.cell(row=r, column=17).number_format = 'R$ #,##0.00'

    econ_l = f'=IF(OR(K{r}="Sem dados",L{r}="",L{r}=0,K{r}>=L{r}),0,H{r}-(G{r}/L{r}))'
    ws_ac.cell(row=r, column=18, value=econ_l)
    ws_ac.cell(row=r, column=18).number_format = '0.0'

    econ_r = f'=R{r}*O{r}'
    ws_ac.cell(row=r, column=19, value=econ_r)
    ws_ac.cell(row=r, column=19).number_format = 'R$ #,##0.00'

    for c in range(1, len(ac_headers) + 1):
        ws_ac.cell(row=r, column=c).border = BORDER
        if ws_ac.cell(row=r, column=c).font.name != FONT_NAME:
            ws_ac.cell(row=r, column=c).font = Font(name=FONT_NAME, size=10)
    r += 1
last_data_row = r - 1

for col, w in zip(range(1, len(ac_headers) + 1),
                   [9, 11, 15, 12, 30, 14, 14, 14, 10, 10, 9, 10, 11, 20, 11, 11, 14, 13, 14]):
    ws_ac.column_dimensions[get_column_letter(col)].width = w

# conditional formatting for Status column and Desvio
ws_ac.conditional_formatting.add(
    f'N{first_data_row}:N{last_data_row}',
    CellIsRule(operator='equal', formula=['"Crítico: abaixo da meta"'], fill=RED_FILL))
ws_ac.conditional_formatting.add(
    f'N{first_data_row}:N{last_data_row}',
    CellIsRule(operator='equal', formula=['"Na meta"'], fill=GREEN_FILL))
ws_ac.conditional_formatting.add(
    f'N{first_data_row}:N{last_data_row}',
    CellIsRule(operator='equal', formula=['"Acima da meta"'], fill=GREEN_FILL))

print("Analise_Consumo rows:", first_data_row, "-", last_data_row)

AC = ws_ac.title
AC_FIRST, AC_LAST = first_data_row, last_data_row

# ============================================================
# 4) PAINEL (dashboard) sheet
# ============================================================
print("Building Painel sheet...")
ws_p = wb.create_sheet('Painel')
ws_p.sheet_view.showGridLines = False
ws_p['B2'] = 'PAINEL EXECUTIVO — GESTÃO DE MÉDIA DE COMBUSTÍVEL'
ws_p['B2'].font = TITLE_FONT
ws_p.merge_cells('B2:J2')
ws_p['B3'] = 'ROTA Transportes - Itabuna | Boletim 08/09/2026 a 14/09/2026 | Frota válida (Base_Frota) presente no boletim'
ws_p['B3'].font = Font(name=FONT_NAME, italic=True, size=10, color='595959')
ws_p.merge_cells('B3:J3')

kpi_defs = [
    ('Veículos válidos analisados', f'=COUNTA({AC}!A{AC_FIRST}:A{AC_LAST})', '0'),
    ('Km totais (todos os dias)', f'=SUM({AC}!F{AC_FIRST}:F{AC_LAST})', '#,##0'),
    ('Km em dias com reporte', f'=SUM({AC}!G{AC_FIRST}:G{AC_LAST})', '#,##0'),
    ('Litros consumidos (dias com reporte)', f'=SUM({AC}!H{AC_FIRST}:H{AC_LAST})', '#,##0'),
    ('Km/L real da frota (km ÷ litros)', f'=IFERROR(B6/B7,"")'.replace('B6', f"SUM({AC}!G{AC_FIRST}:G{AC_LAST})").replace('B7', f"SUM({AC}!H{AC_FIRST}:H{AC_LAST})"), '0.00'),
    ('Custo total diesel no período (R$)', f'=SUM({AC}!Q{AC_FIRST}:Q{AC_LAST})', 'R$ #,##0.00'),
    ('Economia potencial no período (L)', f'=SUM({AC}!R{AC_FIRST}:R{AC_LAST})', '#,##0.0'),
    ('Economia potencial no período (R$)', f'=SUM({AC}!S{AC_FIRST}:S{AC_LAST})', 'R$ #,##0.00'),
]
row = 5
col = 2
for i, (label, formula, fmt) in enumerate(kpi_defs):
    rr = row + (i // 2) * 4
    cc = col + (i % 2) * 4
    box = ws_p.cell(row=rr, column=cc, value=label)
    box.font = Font(name=FONT_NAME, size=10, color='595959')
    ws_p.merge_cells(start_row=rr, start_column=cc, end_row=rr, end_column=cc + 2)
    val = ws_p.cell(row=rr + 1, column=cc, value=formula)
    val.font = Font(name=FONT_NAME, size=20, bold=True, color='1F3864')
    val.number_format = fmt
    ws_p.merge_cells(start_row=rr + 1, start_column=cc, end_row=rr + 1, end_column=cc + 2)
    for rrr in (rr, rr + 1):
        for ccc in range(cc, cc + 3):
            ws_p.cell(row=rrr, column=ccc).fill = SUBHEADER_FILL

r_class = row + 4 * 4 + 1
ws_p.cell(row=r_class, column=2, value='Continuidade do reporte de média (telemetria) — veículos válidos')
ws_p.cell(row=r_class, column=2).font = Font(name=FONT_NAME, bold=True, size=11)
ws_p.merge_cells(start_row=r_class, start_column=2, end_row=r_class, end_column=6)
hh = ['Classificação', 'Qtde de veículos', '% do total']
for i, h in enumerate(hh):
    ws_p.cell(row=r_class + 1, column=2 + i, value=h)
style_header(ws_p, r_class + 1, len(hh), start_col=2)
classif_order = ['Reporte contínuo', 'Reporte retomado / aparente correção', 'Intermitente',
                  'Deixou de reportar', 'Sem reporte em todo o período observado',
                  'Dados insuficientes ou inconsistentes']
rr = r_class + 2
for cl in classif_order:
    ws_p.cell(row=rr, column=2, value=cl)
    ws_p.cell(row=rr, column=3, value=f'=COUNTIF({AC}!E{AC_FIRST}:E{AC_LAST},B{rr})')
    ws_p.cell(row=rr, column=4, value=f'=IFERROR(C{rr}/COUNTA({AC}!A{AC_FIRST}:A{AC_LAST}),0)')
    ws_p.cell(row=rr, column=4).number_format = '0.0%'
    for c in range(2, 5):
        ws_p.cell(row=rr, column=c).border = BORDER
    rr += 1

r_setor = rr + 2
ws_p.cell(row=r_setor, column=2, value='Comparação entre Unidades Operacionais equivalentes (Setor)')
ws_p.cell(row=r_setor, column=2).font = Font(name=FONT_NAME, bold=True, size=11)
ws_p.merge_cells(start_row=r_setor, start_column=2, end_row=r_setor, end_column=7)
hh2 = ['Setor', 'Veículos', 'Km em dias com reporte', 'Litros', 'Km/L real (ponderado)', 'Custo total (R$)']
for i, h in enumerate(hh2):
    ws_p.cell(row=r_setor + 1, column=2 + i, value=h)
style_header(ws_p, r_setor + 1, len(hh2), start_col=2)
setores = sorted(set(f['setor'] for f in frota_rows if f['setor']))
rr = r_setor + 2
for s in setores:
    ws_p.cell(row=rr, column=2, value=s)
    ws_p.cell(row=rr, column=3, value=f'=COUNTIF({AC}!D{AC_FIRST}:D{AC_LAST},B{rr})')
    ws_p.cell(row=rr, column=4, value=f'=SUMIFS({AC}!G{AC_FIRST}:G{AC_LAST},{AC}!D{AC_FIRST}:D{AC_LAST},B{rr})')
    ws_p.cell(row=rr, column=5, value=f'=SUMIFS({AC}!H{AC_FIRST}:H{AC_LAST},{AC}!D{AC_FIRST}:D{AC_LAST},B{rr})')
    ws_p.cell(row=rr, column=6, value=f'=IFERROR(D{rr}/E{rr},"")')
    ws_p.cell(row=rr, column=6).number_format = '0.00'
    ws_p.cell(row=rr, column=7, value=f'=SUMIFS({AC}!Q{AC_FIRST}:Q{AC_LAST},{AC}!D{AC_FIRST}:D{AC_LAST},B{rr})')
    ws_p.cell(row=rr, column=7).number_format = 'R$ #,##0.00'
    for c in range(2, 8):
        ws_p.cell(row=rr, column=c).border = BORDER
    rr += 1

r_rank = rr + 2
ws_p.cell(row=r_rank, column=2,
          value='Ranking de oportunidades — Top 10 veículos por economia potencial estimada (R$)')
ws_p.cell(row=r_rank, column=2).font = Font(name=FONT_NAME, bold=True, size=11)
ws_p.merge_cells(start_row=r_rank, start_column=2, end_row=r_rank, end_column=8)
hh3 = ['#', 'Prefixo', 'Placa', 'Setor', 'Km/L real', 'Meta', 'Economia potencial (R$)', 'Status']
for i, h in enumerate(hh3):
    ws_p.cell(row=r_rank + 1, column=2 + i, value=h)
style_header(ws_p, r_rank + 1, len(hh3), start_col=2)
rr = r_rank + 2
for k in range(1, 11):
    ws_p.cell(row=rr, column=2, value=k)
    large_formula = f'=IFERROR(LARGE({AC}!$S${AC_FIRST}:$S${AC_LAST},{k}),"")'
    match_row = (f'MATCH(1,INDEX(({AC}!$S${AC_FIRST}:$S${AC_LAST}=H{rr})*'
                 f'(COUNTIF($C${r_rank+2}:$C{rr},{AC}!$A${AC_FIRST}:$A${AC_LAST})=0),0),0)')
    ws_p.cell(row=rr, column=3, value=f'=IFERROR(INDEX({AC}!$A${AC_FIRST}:$A${AC_LAST},{match_row}),"")')
    ws_p.cell(row=rr, column=4, value=f'=IFERROR(INDEX({AC}!$B${AC_FIRST}:$B${AC_LAST},{match_row}),"")')
    ws_p.cell(row=rr, column=5, value=f'=IFERROR(INDEX({AC}!$D${AC_FIRST}:$D${AC_LAST},{match_row}),"")')
    ws_p.cell(row=rr, column=6, value=f'=IFERROR(INDEX({AC}!$K${AC_FIRST}:$K${AC_LAST},{match_row}),"")')
    ws_p.cell(row=rr, column=6).number_format = '0.00'
    ws_p.cell(row=rr, column=7, value=f'=IFERROR(INDEX({AC}!$L${AC_FIRST}:$L${AC_LAST},{match_row}),"")')
    ws_p.cell(row=rr, column=7).number_format = '0.00'
    ws_p.cell(row=rr, column=8, value=large_formula)
    ws_p.cell(row=rr, column=8).number_format = 'R$ #,##0.00'
    ws_p.cell(row=rr, column=9, value=f'=IFERROR(INDEX({AC}!$N${AC_FIRST}:$N${AC_LAST},{match_row}),"")')
    for c in range(2, 10):
        ws_p.cell(row=rr, column=c).border = BORDER
    rr += 1

autofit(ws_p, [2] + [14] * 9)
ws_p.sheet_view.showGridLines = False

# ============================================================
# 5) ANÁLISE POR VEÍCULO (rules-based, computed values)
# ============================================================
print("Building Análise por Veículo sheet...")
ws_av = wb.create_sheet('Análise por Veículo')
headers_av = ['Placa', 'Prefixo', 'UO/Localidade', 'Setor', 'Classificação', 'Primeiro dia com registro',
              'Último dia com registro', 'Dias com registro no boletim', 'Dias com média numérica',
              'Dias com tracinho', 'Situação no último registro', 'Distância (km) nos dias com tracinho',
              'Correspondência/ressalva de cruzamento', 'Observação objetiva']
for i, h in enumerate(headers_av):
    ws_av.cell(row=1, column=1 + i, value=h)
style_header(ws_av, 1, len(headers_av))
ws_av.freeze_panes = 'A2'
r = 2
for a in res['analise']:
    vals = [a['placa'], a['prefixo'], a['uo'], a['setor'], a['classificacao'], a['primeiro_dia'],
            a['ultimo_dia'], a['dias_com_registro'], a['dias_numericos'], a['dias_tracinho'],
            a['situacao_ultimo_registro'], a['distancia_dias_tracinho'],
            a['correspondencia_ressalva'], a['observacao']]
    for i, v in enumerate(vals):
        c = ws_av.cell(row=r, column=1 + i, value=v)
        c.font = Font(name=FONT_NAME, size=10)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        c.border = BORDER
    fill = None
    if a['classificacao'] == 'Deixou de reportar':
        fill = RED_FILL
    elif a['classificacao'] == 'Sem reporte em todo o período observado':
        fill = RED_FILL
    elif a['classificacao'] == 'Intermitente':
        fill = YELLOW_FILL2
    elif a['classificacao'] == 'Reporte retomado / aparente correção':
        fill = YELLOW_FILL2
    elif a['classificacao'] == 'Reporte contínuo':
        fill = GREEN_FILL
    if fill:
        ws_av.cell(row=r, column=5).fill = fill
    r += 1
autofit(ws_av, [10, 9, 26, 12, 30, 13, 13, 10, 10, 10, 20, 14, 34, 60])
for c in range(1, len(headers_av) + 1):
    ws_av.column_dimensions[get_column_letter(c)].width = ws_av.column_dimensions[get_column_letter(c)].width
ws_av.row_dimensions[1].height = 30

# ============================================================
# 6) OCORRÊNCIAS
# ============================================================
print("Building Ocorrências sheet...")
ws_oc = wb.create_sheet('Ocorrências')
headers_oc = ['Data', 'Prefixo', 'Placa', 'Média (registro original)', 'Distância (km)', 'Tipo de ocorrência']
for i, h in enumerate(headers_oc):
    ws_oc.cell(row=1, column=1 + i, value=h)
style_header(ws_oc, 1, len(headers_oc))
ws_oc.freeze_panes = 'A2'
r = 2
for o in sorted(res['ocorrencias'], key=lambda x: (x['veiculo_prefixo'] or '', x['data'])):
    vals = [o['data'], o['veiculo_prefixo'], o['veiculo_placa'], o['media'], o['distancia'], o['tipo']]
    for i, v in enumerate(vals):
        c = ws_oc.cell(row=r, column=1 + i, value=v)
        c.font = Font(name=FONT_NAME, size=10)
        c.border = BORDER
    if 'Início' in o['tipo']:
        ws_oc.cell(row=r, column=6).fill = RED_FILL
    else:
        ws_oc.cell(row=r, column=6).fill = GREEN_FILL
    r += 1
autofit(ws_oc, [13, 10, 12, 20, 13, 40])

# ============================================================
# 7) EXCLUÍDOS
# ============================================================
print("Building Excluídos sheet...")
ws_ex = wb.create_sheet('Excluídos')
ws_ex['A1'] = ('Veículos presentes no Boletim que NÃO constam na Base_Frota — mantidos fora da análise '
               'principal de continuidade de reporte, registros originais preservados na aba do boletim.')
ws_ex['A1'].font = Font(name=FONT_NAME, italic=True, size=9, color='595959')
ws_ex.merge_cells('A1:F1')
headers_ex = ['Placa', 'Prefixo', 'UO/Localidade', 'Período dos registros', 'Qtd. de registros no boletim', 'Motivo da exclusão']
for i, h in enumerate(headers_ex):
    ws_ex.cell(row=2, column=1 + i, value=h)
style_header(ws_ex, 2, len(headers_ex))
ws_ex.freeze_panes = 'A3'
r = 3
for e in res['excluidos']:
    vals = [e['placa'], e['prefixo'], e['uo_localidade'], e['periodo_registros'], e['qtd_registros'], e['motivo']]
    for i, v in enumerate(vals):
        c = ws_ex.cell(row=r, column=1 + i, value=v)
        c.font = Font(name=FONT_NAME, size=10)
        c.border = BORDER
    r += 1
autofit(ws_ex, [13, 10, 26, 20, 15, 60])

# ============================================================
# 8) Veículos da Base_Frota sem registro no boletim
# ============================================================
print("Building Sem_Registro_Boletim sheet...")
ws_sb = wb.create_sheet('Sem_Registro_Boletim')
ws_sb['A1'] = 'Veículos da Base_Frota que NÃO possuem nenhum registro no boletim do período (08/09 a 14/09/2026).'
ws_sb['A1'].font = Font(name=FONT_NAME, italic=True, size=9, color='595959')
ws_sb.merge_cells('A1:D1')
headers_sb = ['Placa', 'Prefixo', 'UO', 'Setor']
for i, h in enumerate(headers_sb):
    ws_sb.cell(row=2, column=1 + i, value=h)
style_header(ws_sb, 2, len(headers_sb))
r = 3
for s in res['sem_boletim']:
    vals = [s['placa'], s['prefixo'], s['uo'], s['setor']]
    for i, v in enumerate(vals):
        c = ws_sb.cell(row=r, column=1 + i, value=v)
        c.font = Font(name=FONT_NAME, size=10)
        c.border = BORDER
    r += 1
autofit(ws_sb, [13, 10, 14, 14])

# ============================================================
# 9) RESUMO
# ============================================================
print("Building Resumo sheet...")
ws_r = wb.create_sheet('Resumo')
ws_r.sheet_view.showGridLines = False
ws_r['B2'] = 'RESUMO — DISPONIBILIDADE E CONTINUIDADE DO REPORTE DE MÉDIA DE COMBUSTÍVEL'
ws_r['B2'].font = TITLE_FONT
ws_r.merge_cells('B2:F2')
ws_r['B3'] = 'Boletim ROTA Transportes - Itabuna | Período: 08/09/2026 a 14/09/2026'
ws_r['B3'].font = Font(name=FONT_NAME, italic=True, size=10, color='595959')

resumo = res['resumo']
linhas_resumo = [
    ('Total de veículos na Base_Frota', resumo['total_base_frota']),
    ('Total de veículos válidos presentes no boletim', resumo['total_validos_presentes_boletim']),
    ('Total de veículos válidos sem nenhum registro no boletim', resumo['total_validos_sem_registro_boletim']),
    ('Total de veículos excluídos (boletim, fora da Base_Frota)', resumo['excluidos_total']),
    ('Período esperado', resumo['periodo_esperado']),
    ('Confirmação do período', resumo['periodo_confirmado']),
    ('Registros com placa ilegível/corrompida no boletim', resumo['placas_corrompidas_no_boletim']),
    ('Registros com inconsistência de hodômetro (>5 km de divergência)', resumo['dias_com_problema_hodometro']),
    ('Dias com registros simultâneos numérico e tracinho (mesmo veículo/dia)', resumo['dias_mistos_detectados']),
]
r = 5
for label, val in linhas_resumo:
    ws_r.cell(row=r, column=2, value=label).font = Font(name=FONT_NAME, bold=True, size=10)
    ws_r.cell(row=r, column=5, value=val).font = Font(name=FONT_NAME, size=10)
    r += 1

r += 1
ws_r.cell(row=r, column=2, value='Classificação (base: veículos válidos presentes no boletim)').font = Font(name=FONT_NAME, bold=True, size=11)
r += 1
hh = ['Classificação', 'Quantidade', '% do total']
for i, h in enumerate(hh):
    ws_r.cell(row=r, column=2 + i, value=h)
style_header(ws_r, r, len(hh), start_col=2)
r += 1
for cl in classif_order:
    qtd = resumo['classificacao_contagem'].get(cl, 0)
    pct = resumo['classificacao_percentual'].get(cl, 0)
    ws_r.cell(row=r, column=2, value=cl)
    ws_r.cell(row=r, column=3, value=qtd)
    ws_r.cell(row=r, column=4, value=pct / 100)
    ws_r.cell(row=r, column=4).number_format = '0.0%'
    for c in range(2, 5):
        ws_r.cell(row=r, column=c).border = BORDER
    r += 1

r += 1
ws_r.cell(row=r, column=2, value='Excluídos por localidade (UO informada no boletim)').font = Font(name=FONT_NAME, bold=True, size=11)
r += 1
for h_i, h in enumerate(['Localidade', 'Quantidade']):
    ws_r.cell(row=r, column=2 + h_i, value=h)
style_header(ws_r, r, 2, start_col=2)
r += 1
for loc, qtd in resumo['excluidos_por_localidade'].items():
    ws_r.cell(row=r, column=2, value=loc)
    ws_r.cell(row=r, column=3, value=qtd)
    for c in range(2, 4):
        ws_r.cell(row=r, column=c).border = BORDER
    r += 1

r += 1
ws_r.cell(row=r, column=2,
          value='Lista prioritária de investigação — veículos que rodaram sem reportar média '
                '(ordenado por classificação de risco e distância percorrida sem reporte)').font = Font(name=FONT_NAME, bold=True, size=11)
ws_r.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
r += 1
hh4 = ['Prefixo', 'Placa', 'Classificação', 'Situação no último registro', 'Km nos dias com tracinho', 'Último dia com registro']
for i, h in enumerate(hh4):
    ws_r.cell(row=r, column=2 + i, value=h)
style_header(ws_r, r, len(hh4), start_col=2)
r += 1
for p in res['prioridade_investigacao']:
    vals = [p['prefixo'], p['placa'], p['classificacao'], p['situacao_ultimo_registro'],
            p['distancia_dias_tracinho'], p['ultimo_dia']]
    for i, v in enumerate(vals):
        c = ws_r.cell(row=r, column=2 + i, value=v)
        c.font = Font(name=FONT_NAME, size=10)
        c.border = BORDER
    r += 1

r += 2
ws_r.cell(row=r, column=2, value='Regras e limitações desta análise').font = Font(name=FONT_NAME, bold=True, size=11)
r += 1
regras_texto = [
    '• "-" (tracinho) na Média de Consumo = telemetria não importou a média naquele registro (sem comunicação/sem dado).',
    '• Qualquer valor numérico, inclusive 0,00, conta como reporte (comunicação ocorreu).',
    '• Datas sem linha no boletim NÃO foram tratadas como interrupção — apenas a sequência de dias efetivamente registrados foi avaliada.',
    '• Retomada de números após tracinho é tratada como indício de correção, sem confirmação de manutenção (não há registro de OS nesta base).',
    '• Não foi atribuída causa mecânica nem responsabilidade ao motorista em nenhuma classificação.',
    '• km/L é sempre calculado por (km totais ÷ litros totais) no grupo/veículo — nunca pela média simples das médias diárias.',
    '• Frota válida = exclusivamente veículos cadastrados na aba Base_Frota, cruzados pela Placa (normalizada) e, em caso de divergência, pelo Prefixo.',
]
for t in regras_texto:
    ws_r.cell(row=r, column=2, value=t).font = Font(name=FONT_NAME, size=9)
    ws_r.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    ws_r.cell(row=r, column=2).alignment = Alignment(wrap_text=True)
    r += 1

autofit(ws_r, [3, 30, 30, 22, 18, 16])

wb.save(OUT)
print("Saved:", OUT)
